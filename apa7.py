# apa7.py: aplica formato APA 7 a un .docx sin cambiar la información
# Uso: python apa7.py entrada.docx salida.docx
import os
import re
import sys
import unicodedata

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FUENTE = "Arial"
TAM = Pt(11)  # APA 7: Arial va en 11 pt
SANGRIA = Cm(1.27)
QUITAR_INDICE = False  # APA 7 no exige índice: True lo elimina (si el docente no lo pide)

# Palabras que van en minúscula al pasar un texto en MAYÚSCULAS a tipo título
CONECTORES = {"de", "del", "la", "las", "el", "los", "y", "e", "o", "u", "en", "a",
              "para", "por", "con", "sin", "al"}

# Tildes que faltan en los títulos escritos en MAYÚSCULAS
TILDES = {"delimitacion": "delimitación", "justificacion": "justificación",
          "metodologia": "metodología", "capitulo": "capítulo", "titulo": "título",
          "indice": "índice"}

ENLACE_MD = re.compile(r"\[([^\]]*)\]\((https?://[^)\s]*)\)?")  # [texto](url)
ES_REFERENCIA = re.compile(r"\((s\. ?f\.|\d{4})[a-z]?[,)]")      # (2024) / (s. f.)


# ----------------------------------------------------------------- utilidades
def fmt_parrafo(pf, sangria=None, colgante=False):
    """Interlineado doble, sin espacio extra y sangría APA."""
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if colgante:
        pf.left_indent = SANGRIA
        pf.first_line_indent = -SANGRIA
    elif sangria is not None:
        pf.left_indent = Cm(0)
        pf.first_line_indent = sangria


def quitar(elemento, *etiquetas):
    """Elimina hijos de pPr como numeración automática o tabuladores."""
    ppr = elemento.find(qn("w:pPr"))
    if ppr is None:
        return
    for etiqueta in etiquetas:
        for hijo in ppr.findall(qn(etiqueta)):
            ppr.remove(hijo)


def borrar(p):
    p._p.getparent().remove(p._p)


def reemplazar_texto(p, nuevo):
    """Pone 'nuevo' en el primer run y vacía los demás (conserva el formato del primero)."""
    runs = p.runs
    if not runs:
        return
    runs[0].text = nuevo
    for r in runs[1:]:
        r.text = ""


def quitar_espaciado_automatico(raiz):
    """El 'espaciado automático' de Word ignora el 0 pt: se desactiva."""
    for sp in raiz.iter(qn("w:spacing")):
        for attr in ("w:beforeAutospacing", "w:afterAutospacing"):
            sp.attrib.pop(qn(attr), None)


def desvincular_numeracion(doc, ids_estilo):
    """Quita la numeración automática que Word asocia a los estilos de título."""
    try:
        numbering = doc.part.numbering_part.element
    except (KeyError, NotImplementedError):
        return
    for lvl in numbering.iter(qn("w:lvl")):
        ps = lvl.find(qn("w:pStyle"))
        if ps is not None and ps.get(qn("w:val")) in ids_estilo:
            lvl.remove(ps)


def a_oracion(texto):
    """METODOLOGIA -> Metodología; respeta números romanos (CAPITULO II)."""
    palabras = []
    for i, w in enumerate(texto.split(" ")):
        if re.fullmatch(r"[IVXLC]+", w):
            palabras.append(w)
            continue
        w = w.lower()
        w = TILDES.get(w, w)
        palabras.append(w[:1].upper() + w[1:] if i == 0 else w)
    return " ".join(palabras)


def a_tipo_titulo(texto):
    """FACULTAD DE CIENCIAS -> Facultad de Ciencias (para nombres en la portada).
    Respeta siglas seguidas de número (INV 400) y números romanos."""
    palabras = texto.split(" ")
    salida = []
    for i, w in enumerate(palabras):
        siguiente = palabras[i + 1] if i + 1 < len(palabras) else ""
        if re.fullmatch(r"[IVXLC]+", w) or re.match(r"^\d", siguiente):
            salida.append(w)
            continue
        base = w.lower()
        if i > 0 and base in CONECTORES:
            salida.append(base)
        else:
            salida.append(base[:1].upper() + base[1:])
    return " ".join(salida)


def quitar_numeracion(texto):
    """'2.1.1. Concepto' -> 'Concepto' (APA 7 no numera los títulos)."""
    return re.sub(r"^\d+(\.\d+)*\.?\s+", "", texto)


def fmt_estilo(estilo, negrita=False, cursiva=False, centrado=False):
    """Formatea un estilo de título según APA."""
    estilo.base_style = None  # que no herede de "Párrafo de lista"
    quitar(estilo.element, "w:numPr", "w:tabs", "w:ind")
    f = estilo.font
    f.name = FUENTE
    f.size = TAM
    f.bold = negrita
    f.italic = cursiva
    f.color.rgb = RGBColor(0, 0, 0)
    pf = estilo.paragraph_format
    fmt_parrafo(pf, sangria=Cm(0))
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrado else WD_ALIGN_PARAGRAPH.LEFT
    pf.keep_with_next = True


def numero_de_pagina(seccion):
    """Número de página arriba a la derecha."""
    p = seccion.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.text = texto
        run._r.append(el)
    run.font.name = FUENTE
    run.font.size = TAM


def nivel_titulo(p):
    """1, 2 o 3 si el párrafo es un título; 0 si es texto normal."""
    t = p.text.strip().lstrip("#").strip()
    if not t or len(t) > 120:
        return 0
    if re.match(r"^\d+\.\d+\.\d+\.?\s", t):
        return 3  # 2.1.1. ...
    if re.match(r"^\d+\.\d+\.?\s", t):
        return 2  # 2.1. ...
    if p.style.name in ("Heading 1", "Heading 2", "Heading 3"):
        return int(p.style.name[-1])
    if t.isupper() and len(t) < 60:
        return 1  # CAPITULO II
    return 0


def es_inicio_referencias(p):
    t = p.text.strip().lstrip("#").strip().lower()
    return t.startswith("referencias") or t.startswith("fuentes bibliogr")


def es_vacio(p):
    """Párrafo sin texto, sin imagen y sin salto de sección."""
    if p.text.strip():
        return False
    if p._p.xpath(".//w:drawing") or p._p.xpath("./w:pPr/w:sectPr"):
        return False
    return True


# ----------------------------------------------------------------- portada e índice
def espaciado_original(doc):
    """Espaciado que tenía el estilo Normal antes de cambiarlo (para la portada)."""
    atributos = {}
    defecto = doc.styles.element.find(qn("w:docDefaults"))
    normal = doc.styles["Normal"].element
    for origen in (defecto, normal):
        if origen is None:
            continue
        for sp in origen.iter(qn("w:spacing")):
            atributos.update(sp.attrib)
    return atributos


def fmt_portada(parrafos, espaciado):
    """La portada conserva su diseño original: solo cambia la fuente a Arial."""
    for p in parrafos:
        ppr = p._p.get_or_add_pPr()
        sp = ppr.find(qn("w:spacing"))
        if sp is None:
            sp = OxmlElement("w:spacing")
            ppr.append(sp)
        for k, v in espaciado.items():
            if sp.get(k) is None:
                sp.set(k, v)
        if p._p.xpath("./w:pPr/w:ind") == [] and p.style.name == "Normal":
            p.paragraph_format.first_line_indent = Cm(0)
        texto = p.text.strip()
        if texto.isupper():  # APA 7 no usa MAYÚSCULAS sostenidas
            reemplazar_texto(p, a_tipo_titulo(texto))
        for r in p.runs:
            r.font.name = FUENTE


def fmt_indice(doc):
    """Índice en hoja nueva, estilos legibles y actualización al abrir en Word."""
    for sdt in doc.element.body.iter(qn("w:sdt")):
        if QUITAR_INDICE:
            sdt.getparent().remove(sdt)
            return
        primero = next(sdt.iter(qn("w:p")), None)
        if primero is not None:
            textos = list(primero.iter(qn("w:t")))
            titulo = "".join(t.text or "" for t in textos).strip()
            if titulo.isupper() and textos:  # INDICE -> Índice
                textos[0].text = a_oracion(titulo)
                for t in textos[1:]:
                    t.text = ""
            ppr = primero.get_or_add_pPr()
            if ppr.find(qn("w:pageBreakBefore")) is None:
                ppr.insert(0, OxmlElement("w:pageBreakBefore"))
        # El último renglón del índice (vacío, cierra el campo) se hace mínimo
        # para que no se pase solo a una hoja en blanco
        ultimo = list(sdt.iter(qn("w:p")))[-1]
        if not "".join(t.text or "" for t in ultimo.iter(qn("w:t"))).strip():
            ppr = ultimo.get_or_add_pPr()
            for viejo in ppr.findall(qn("w:spacing")):
                ppr.remove(viejo)
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:before"), "0")
            sp.set(qn("w:after"), "0")
            sp.set(qn("w:line"), "20")
            sp.set(qn("w:lineRule"), "exact")
            ppr.append(sp)
        break
    for n in (1, 2, 3):
        try:
            estilo = doc.styles[f"TOC {n}"]
        except KeyError:
            continue
        estilo.font.name = FUENTE
        estilo.font.size = TAM
        pf = estilo.paragraph_format
        pf.line_spacing = 1.15
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        pf.left_indent = Cm(0.6 * (n - 1))
        pf.first_line_indent = Cm(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT


def pedir_actualizacion_al_abrir(ruta):
    """Plan B: Word preguntará '¿Actualizar los campos?' al abrir el archivo."""
    doc = Document(ruta)
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.append(uf)
    doc.save(ruta)


def actualizar_indice_con_word(ruta):
    """Abre el archivo en Word (sin mostrarlo), regenera el índice y lo guarda."""
    try:
        import win32com.client
    except ImportError:
        print("Aviso: falta pywin32 (pip install pywin32). Actualiza el índice en Word a mano.")
        pedir_actualizacion_al_abrir(ruta)
        return
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        documento = word.Documents.Open(os.path.abspath(ruta))
        try:  # estilos del índice y su título: misma letra, sin sangría ni doble espacio
            for nivel, id_estilo in enumerate((-20, -21, -22)):  # TOC 1, 2 y 3 en Word
                estilo = documento.Styles(id_estilo)
                estilo.Font.Name = FUENTE
                estilo.Font.Size = TAM.pt
                formato = estilo.ParagraphFormat
                formato.FirstLineIndent = 0
                formato.LeftIndent = 17 * nivel
                formato.LineSpacingRule = 0  # sencillo
                formato.SpaceBefore = 0
                formato.SpaceAfter = 6
            for parrafo in documento.Paragraphs:
                if parrafo.Range.Text.strip().lower() in ("índice", "indice"):
                    letra = parrafo.Range.Font
                    letra.Name = FUENTE
                    letra.Size = TAM.pt
                    letra.Bold = True
                    letra.Color = 0
                    break
        except Exception as error:
            print("Aviso: no se pudo dar formato al índice:", error)
        for indice in documento.TablesOfContents:
            indice.Update()  # = "Actualizar toda la tabla"
        documento.Save()
        documento.Close()
        print("Índice actualizado y guardado.")
    except Exception as error:
        print("Aviso: no se pudo actualizar el índice con Word:", error)
        pedir_actualizacion_al_abrir(ruta)
    finally:
        word.Quit()


# ----------------------------------------------------------------- referencias
def limpiar_referencia(t):
    t = ENLACE_MD.sub(lambda m: m.group(2), t)  # [dhl](https://...) -> https://...
    t = re.sub(r"\s+", " ", t).strip()
    return t


def extraer_referencias(parrafos):
    """Saca del texto todos los bloques 'Referencias' y devuelve sus entradas."""
    entradas, notas = [], []
    en_bloque = False
    for p in parrafos:
        if es_inicio_referencias(p):
            en_bloque = True
            borrar(p)
            continue
        if not en_bloque:
            continue
        if nivel_titulo(p):  # siguiente título: termina el bloque
            en_bloque = False
            continue
        t = p.text.strip()
        borrar(p)
        if not t or t.startswith("#"):
            continue  # vacío o subtítulo "### Título de la obra"
        t = limpiar_referencia(t)
        if re.match(r"^https?://", t) and entradas:
            entradas[-1] += " " + t  # URL que quedó en otro renglón
        elif ES_REFERENCIA.search(t):
            entradas.append(t)
        elif re.match(r"^\S*https?://", t) and entradas:
            entradas[-1] += " " + t
        else:
            notas.append(t)  # nota del texto original, no es una referencia
            urls = re.findall(r"https?://\S+", t)
            if urls and entradas:  # si la nota traía el enlace de la referencia, se conserva
                entradas[-1] += " " + " ".join(urls)
    return entradas, notas


FECHA = re.compile(r"\((s\. ?f\.|\d{4})[^)]*\)\. ")  # (2024, 1 de mayo). / (s. f.).


def clave_orden(t):
    t = unicodedata.normalize("NFKD", t.replace("*", "")).encode("ascii", "ignore").decode()
    return t.lower().replace("(s. f.", "(0000").replace("(s.f.", "(0000")  # s. f. va primero


def clave_apa(t):
    """Orden APA: autor, luego año (s. f. primero) y, con el mismo año, por título."""
    m = FECHA.search(t)
    if not m:
        return (clave_orden(t), "", "")
    anio = "0000" if m.group(1).startswith("s") else m.group(1)
    return (clave_orden(t[:m.start()]), anio, clave_orden(t[m.end():]))


def deduplicar(entradas):
    vistas, unicas = set(), []
    for e in entradas:
        k = re.sub(r"\W", "", clave_orden(e))
        if k not in vistas:
            vistas.add(k)
            unicas.append(e)
    return unicas


def marcar_cursiva(t):
    """Si la referencia no trae *cursiva*, la deduce según APA 7."""
    if "*" in t:
        return t
    # Artículo de revista: cursiva en "Nombre de la revista, volumen"
    m = re.search(r"\. ([^.]+?, \d+)(\(\d+\))?, (Article|\d)", t)
    if m:
        return t[:m.start(1)] + "*" + m.group(1) + "*" + t[m.end(1):]
    # Libro, informe o página web: cursiva en el título (tras la fecha)
    m = FECHA.search(t)
    if not m:
        return t
    ini = m.end()
    resto = t[ini:]
    fin = len(resto)
    for patron in (r"(?<! vs)\. ", r" \[", r" \((\d+(st|nd|rd|th)|Rev\.|Boletín|\d{4} release)"):
        f = re.search(patron, resto)
        if f and f.start() < fin:
            fin = f.start()
    if resto[:fin].startswith("http"):
        return t
    return t[:ini] + "*" + resto[:fin].rstrip(".") + "*" + resto[fin:] if fin else t


def agregar_referencias(doc, entradas):
    titulo = doc.add_paragraph("Referencias", style="Heading 1")
    titulo.paragraph_format.page_break_before = True
    for e in sorted(entradas, key=clave_apa):
        p = doc.add_paragraph()
        fmt_parrafo(p.paragraph_format, colgante=True)
        for i, trozo in enumerate(marcar_cursiva(e).split("*")):
            if trozo:
                r = p.add_run(trozo)
                r.italic = i % 2 == 1
                r.font.name = FUENTE
                r.font.size = TAM


# ----------------------------------------------------------------- cuerpo
def unir_renglones_cortados(parrafos):
    """Une párrafos partidos a mitad de oración (el texto no cambia)."""
    anterior = None
    for p in parrafos:
        t = p.text
        if anterior is not None and t.strip() and t.strip()[0].islower() \
                and not re.search(r"[.:;?!)]\s*$", anterior.text) \
                and not p._p.xpath("./w:pPr/w:numPr") and not nivel_titulo(p):
            base = anterior.text.rstrip()
            reemplazar_texto(anterior, base + " " + t.strip())
            borrar(p)
            continue
        anterior = p if t.strip() and not nivel_titulo(p) else None


def limpiar_markdown(p):
    """Quita restos de Markdown del texto: [kvk](https://...) y **negrita**."""
    t = p.text
    nuevo = ENLACE_MD.sub("", t).replace("**", "")
    nuevo = re.sub(r"[ \t]{2,}", " ", nuevo).rstrip()
    if nuevo != t.rstrip():
        reemplazar_texto(p, nuevo)


# ----------------------------------------------------------------- principal
def main(entrada, salida):
    doc = Document(entrada)
    espaciado_portada = espaciado_original(doc)

    for s in doc.sections:  # márgenes de 2,54 cm y número de página
        s.top_margin = s.bottom_margin = Cm(2.54)
        s.left_margin = s.right_margin = Cm(2.54)
        numero_de_pagina(s)

    normal = doc.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = TAM
    fmt_parrafo(normal.paragraph_format, sangria=SANGRIA)
    titulos = [doc.styles[f"Heading {n}"] for n in (1, 2, 3)]
    desvincular_numeracion(doc, {e.style_id for e in titulos})
    fmt_estilo(doc.styles["Heading 1"], negrita=True, centrado=True)
    fmt_estilo(doc.styles["Heading 2"], negrita=True)
    fmt_estilo(doc.styles["Heading 3"], negrita=True, cursiva=True)

    # Portada = todo lo anterior al primer título real
    parrafos = list(doc.paragraphs)
    inicio = next(i for i, p in enumerate(parrafos) if p.style.name.startswith("Heading"))
    for p in parrafos[:inicio]:  # renglones vacíos entre el índice y el texto
        if es_vacio(p) and p._p.xpath("preceding-sibling::w:sdt"):
            borrar(p)
    parrafos = list(doc.paragraphs)
    inicio = next(i for i, p in enumerate(parrafos) if p.style.name.startswith("Heading"))
    fmt_portada(parrafos[:inicio], espaciado_portada)
    fmt_indice(doc)
    cuerpo = parrafos[inicio:]
    cuerpo[0].paragraph_format.page_break_before = True

    entradas, notas = extraer_referencias(cuerpo)
    cuerpo = [p for p in doc.paragraphs[inicio:]]
    for p in cuerpo:
        if es_vacio(p):
            borrar(p)
    cuerpo = [p for p in doc.paragraphs[inicio:]]
    unir_renglones_cortados(cuerpo)

    for p in doc.paragraphs[inicio:]:
        nivel = nivel_titulo(p)
        if nivel:
            texto = p.text.strip()
            nuevo = quitar_numeracion(texto.lstrip("#").strip())
            if nuevo.isupper():
                nuevo = a_oracion(nuevo)
            if nuevo != texto:
                reemplazar_texto(p, nuevo)
            for r in p.runs:  # el estilo manda: quitar formato directo
                r.font.bold = r.font.italic = r.font.size = r.font.name = None
            p.style = doc.styles[f"Heading {nivel}"]
            quitar(p._p, "w:numPr", "w:tabs", "w:ind", "w:jc", "w:spacing")
            if nuevo.lower().startswith("capítulo"):  # cada capítulo en hoja nueva
                p.paragraph_format.page_break_before = True
            continue

        limpiar_markdown(p)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if p._p.xpath("./w:pPr/w:numPr"):  # viñetas y listas: conservan su sangría
            fmt_parrafo(p.paragraph_format)
        else:
            fmt_parrafo(p.paragraph_format, sangria=SANGRIA)
        for r in p.runs:
            r.font.name = FUENTE
            r.font.size = TAM

    unicas = deduplicar(entradas)
    agregar_referencias(doc, unicas)

    quitar_espaciado_automatico(doc.element)
    quitar_espaciado_automatico(doc.styles.element)
    doc.save(salida)
    actualizar_indice_con_word(salida)

    print("Listo:", salida)
    print(f"Referencias: {len(entradas)} encontradas, {len(unicas)} sin repetir.")
    if notas:
        print("Notas quitadas de los bloques de referencias (no eran referencias):")
        for n in notas:
            print("  -", n[:100])


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Uso: python apa7.py entrada.docx salida.docx")
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])